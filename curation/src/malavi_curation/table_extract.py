"""Extract structured rows from data tables (.xlsx/.xls/.csv/.tsv/.docx, and PDF).

The host x lineage x locality matrix that MalAvi curates row-by-row lives in a
paper's supplementary table when there is one, and in a printed table in the PDF
when there is not. This module reads either and maps its columns to
MalAvi-relevant fields (lineage, host, country, site, accession, parasite genus,
numbers tested/found) using header synonyms, emitting one dict per data row plus
the flattened text (so the accession miner can still sweep it).

Column detection is heuristic and recall-oriented; a curator confirms. Spreadsheet
support needs the optional 'tables' extra (openpyxl, python-docx); CSV/TSV use the
standard library only.

A supplementary *file* is not a single table. A workbook has several sheets and a
Word document several tables, each with its own header, and only some of them hold
records at all. So the file is read into a list of **blocks** -- one per sheet, per
Word table, or the whole CSV -- and each block is parsed independently. The three
behaviours that follow from that are all things the ground-truth benchmark caught
(2026-07-28) when every block was flattened into one matrix instead:

* **Every table is reachable.** Header detection scans the top of each block, so a
  document's second and later tables are parsed. Previously Perrin et al 2026's
  Table S2 -- the entire host x parasite matrix -- was invisible because Table S1
  came first and the header search never got past it.
* **Data dictionaries are skipped.** A "Dataset S1 metadata" sheet listing
  ``Label | Contents`` is documentation, not data; parsing it produced records
  such as ``DEPARTMENT x IN PERU`` from McNew et al 2021.
* **Compound headers resolve to the right field.** See ``_rank_columns``.

Two further behaviours, both added 2026-07-29 for the same benchmark:

* **PDF tables go down this same path.** ``extract_pdf_tables`` feeds the tables
  carved out of a PDF's layout text (``pdf_extract.carve_layout_tables``) into
  the block parser below. Six of the ten ground-truth papers ship no
  supplementary spreadsheet at all, so before this their record recall could only
  ever be zero.
* **Wide "slot" layouts are expanded.** A row that records a host's co-infections
  across repeated ``Lineage.1..5.Name`` column groups becomes one record per
  occupied slot rather than one record for slot 1. See ``_slot_of``.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

# Field -> header synonyms (matched case-insensitively as substrings).
_COLUMN_SYNONYMS: Dict[str, List[str]] = {
    "lineage_name": ["lineage", "haplotype", "cytb name", "cyt-b name", "parasite lineage", "malavi"],
    # The nucleotide sequence itself, when a supplement prints it. This is what
    # lets a paper's private haplotype namespace ("T001", "G-016") be resolved to
    # the MalAvi lineage it already is -- see lineage_resolve. Guarded by
    # _VALUE_SHAPE, which requires an actual run of nucleotides, so a column
    # merely *named* "sequence" (a sequence count, a sequencing platform) cannot
    # land here.
    "sequence": ["sequence", "nucleotide sequence", "cytb sequence", "dna sequence",
                 "consensus sequence"],
    # "taxon" is deliberately the least specific host synonym: datasets that name
    # the host column "Taxon" (McNew et al 2021) are common, but so are
    # parasite-centric tables, so it is only ever accepted when the column's
    # values actually look like binomials -- see _VALUE_SHAPE.
    # "taxon" and a bare "species" are deliberately the least specific host
    # synonyms: tables that name the host column "Taxon" (McNew et al 2021) or
    # "Species" (Bukauskaitė et al 2024) are common. They are safe only in
    # combination with two guards -- _HOST_EXCLUSION below, which refuses any
    # column explicitly labelled as parasite or vector, and _VALUE_SHAPE, which
    # requires the column's values to actually be binomials.
    "host_species": ["host species", "host", "avian host", "bird species",
                     "host name", "taxon", "species"],
    # The vector table is a MalAvi table in its own right (lineage x vector
    # species x country), and a supplement that carries one usually labels it
    # plainly. Every synonym here also appears in _HOST_EXCLUSION, so a column
    # claimed as a vector can never simultaneously be read as an avian host.
    "vector_species": ["vector species", "vector", "mosquito species", "mosquito",
                       "insect species", "biting midge", "midge", "culicoides",
                       "arthropod"],
    "country": ["country", "locality country", "nation"],
    "site": ["site", "locality", "location", "region", "sampling site", "place"],
    "accession": ["genbank", "accession", "acc. no", "acc no", "acc.number", "accession number", "acc."],
    "parasite_genus": ["parasite genus", "genus", "parasite"],
    # Prevalence. MalAvi stores NUMBER_TESTED / NUMBER_FOUND per record, and a
    # paper that reports prevalence nearly always prints these two counts side by
    # side. Both are guarded by _VALUE_SHAPE: a column only becomes a count if it
    # actually holds integers, which is what stops a "% positive" or "Fisher's
    # exact test" column from landing here.
    "number_tested": ["number tested", "no tested", "n tested", "tested",
                      "number examined", "no examined", "n examined", "examined",
                      "number screened", "screened", "sample size"],
    "number_found": ["number found", "no found", "n found",
                     "number positive", "no positive", "n positive", "positive",
                     "number infected", "no infected", "n infected", "infected"],
}

# A MalAvi record is an *association*: this lineage, in this host. A block that
# cannot supply both halves has no records in it, whatever else it contains, and
# a row missing either half is not a record either.
#
# This is the rule that keeps prose out. `pdftotext -layout` renders a
# two-column journal page as lines with a wide gap down the middle, which is
# structurally indistinguishable from a two-column table; carving PDFs therefore
# proposes many blocks that are really body text. Left to a weaker test they
# produced rows like ``host_species = "eles, Coquillettidia, Lutzia and
# Orthopodomyia. Using the"`` (Schmid et al 2017a) -- which is precisely the
# manufacturing-from-prose failure that commit 4eeb185 removed, arriving by a new
# door. Requiring the association shuts it: prose has no lineage column.
_REQUIRED_RECORD_FIELDS = ("lineage_name", "host_species")

# The same rule for the vector table: a vector record is a lineage found in a
# named vector species. Perrin et al 2026's Supplementary Table S4 is exactly
# this shape (parasite genus | lineage name | mosquito species | frequency |
# GenBank number) and is the reason the vector layer is parsed at all.
_REQUIRED_VECTOR_FIELDS = ("lineage_name", "vector_species")

# Fields that hold a *name*, and so can be spoiled by a category word in the cell
# (see _RESERVED_VALUES). A count or a sequence cannot.
_NAMED_FIELDS = ("lineage_name", "host_species", "vector_species", "country", "site")

# Cell values that mean "empty". Publishers write absence in a dozen ways, and a
# literal "None"/"NA" flowing into lineage_name becomes a false association.
#
# The lineage pattern accepts any compact alphanumeric token, so every spelling of
# "no value" that a publisher invents passes it: "ND" (not determined), "N.D.",
# "TBD" and "Unk" were all accepted as lineage names until they were listed here.
# Spreadsheet error values belong here too -- Fecchio et al 2023b's supplement
# carries "#REF!" cells, which produced two records naming a lineage "#REF!".
_EMPTY_VALUES = {"", "none", "na", "n/a", "n.a.", "null", "nan", "-", "--",
                 "–", "—", ".", "?", "??",
                 # explicit "missing" spellings
                 "nd", "n.d.", "n. d.", "tbd", "unk", "unknown", "unkn",
                 "missing", "not determined", "not tested", "not applicable",
                 "no data", "nodata", "n/d", "--", "---",
                 # spreadsheet error values
                 "#ref!", "#n/a", "#value!", "#name?", "#div/0!", "#null!",
                 "#num!", "#error!", "err:502"}

# Values that are a *category*, not a name, and must never become a lineage or a
# host. "MIX" (a mixed/co-infection) is the one that reached the benchmark: it
# produced the single false record from Bukauskaitė et al 2024. Unlike
# _EMPTY_VALUES these are meaningful cells -- they just do not name anything MalAvi
# can store, so the row is not an association.
_RESERVED_VALUES = {"mix", "mixed", "mixed infection", "coinfection",
                    "co-infection", "multiple", "several", "various",
                    "short", "partial", "new", "novel", "undetermined",
                    "negative", "positive", "pos", "neg", "total", "sum",
                    "subtotal", "other", "others", "unidentified", "sp", "spp"}

# MalAvi lineage names carry no genus prefix (``BUBT3``), but papers print them
# with one (``lBUBT3``, ``hCCF3``, ``pGRW04``) — the leading lowercase letter is
# the parasite genus: l = Leucocytozoon, h = Haemoproteus, p = Plasmodium.
#
# Stripping it is unambiguous, not a guess: checked against the 2026-03-23
# release, **no** MalAvi lineage name begins with a lowercase letter, and none
# matches this pattern at all. Without the strip, every lineage read out of a
# printed table (Harl et al 2026's Table 1 prints all sixteen with the prefix)
# fails to join to MalAvi.
_GENUS_PREFIXED_LINEAGE = re.compile(r"^[hlp](?=[A-Z])")

# A column whose header says it holds a parasite or a vector is never the avian
# host, however well it matches a host synonym. Without this, the weak synonym
# "species" would capture "Parasite species" -- and no value check could catch
# it, because a parasite binomial looks exactly like a host binomial.
_HOST_EXCLUSION = re.compile(
    r"parasite|haemoproteus|plasmodium|leucocytozoon|haemosporidian|"
    r"vector|mosquito|culicoides|midge|insect|arthropod", re.IGNORECASE)

# Block names that mark documentation rather than data. Journals name these
# sheets consistently enough for this to be reliable, and the cost of a wrong
# skip is bounded: the block's text is still swept for accessions, and a curator
# reviews the result either way.
_METADATA_BLOCK = re.compile(
    r"\b(metadata|meta-data|legend|readme|read me|dictionary|glossary|"
    r"column descriptions?|variable descriptions?|key to )\b", re.IGNORECASE)


@dataclass
class ExtractedRows:
    """Structured rows from one supplementary table file."""

    source: Path
    rows: List[Dict[str, Optional[str]]] = field(default_factory=list)
    # Vector associations (lineage x vector species), which MalAvi keeps in its
    # own table and which a supplement often reports in a sheet of its own.
    vectors: List[Dict[str, Optional[str]]] = field(default_factory=list)
    # Accession values read straight out of an accession column. Kept separately
    # from the records because they survive a row that is not an association.
    accessions: List[str] = field(default_factory=list)
    text: str = ""                       # flattened cell text for accession mining
    columns_detected: Dict[str, str] = field(default_factory=dict)  # field -> matched header
    # Per-block detail: what each sheet/table contributed, and why one was skipped.
    blocks: List[Dict[str, object]] = field(default_factory=list)
    # How many identical rows were collapsed (see _deduplicate).
    n_duplicate_rows: int = 0

    def is_empty(self) -> bool:
        return not self.rows and not self.vectors


def _clean(s: Optional[str]) -> str:
    """Lower-case a header cell and normalize its separators to single spaces.

    Publisher headers use dots and underscores as word separators
    ("Lineage.1.Name", "host_species"), so those become spaces here; otherwise a
    synonym like "host species" could never match "host_species".
    """
    text = re.sub(r"[._/\\-]+", " ", (s or "").strip().lower())
    return re.sub(r"\s+", " ", text).strip()


def _clean_value(raw: Optional[object]) -> Optional[str]:
    """Normalize one data cell, mapping every spelling of "absent" to None."""
    if raw is None:
        return None
    text = str(raw).strip()
    return None if text.lower() in _EMPTY_VALUES else text


def _is_reserved(value: Optional[str]) -> bool:
    """Is this cell a category word rather than a name? (see _RESERVED_VALUES)"""
    return bool(value) and value.strip().lower() in _RESERVED_VALUES


def normalize_lineage(value: Optional[str]) -> Optional[str]:
    """Strip a printed lineage's genus prefix, giving the MalAvi form.

    ``lBUBT3`` -> ``BUBT3``; ``CYACYA05`` is returned unchanged.
    """
    if not value:
        return value
    return _GENUS_PREFIXED_LINEAGE.sub("", value.strip())


# One MalAvi lineage name, with or without its printed genus prefix: an optional
# lowercase genus letter, then the name itself (letters then digits, as in
# ``BUBT3``, ``BUTBUT03``, ``ACCGEN01``, ``TURDUS1``, ``GRW04``).
_LINEAGE_TOKEN = re.compile(r"^[hplHPL]?[A-Z]{2,10}[0-9]{1,3}$")

# Delimiters a supplement uses between the lineages of a co-infection. A lineage
# name itself never contains any of them.
_COINFECTION_SPLIT = re.compile(r"\s*(?:,|;|/|\+|&|\band\b)\s*", re.IGNORECASE)


def split_coinfection(value: Optional[str]) -> List[str]:
    """Every lineage named in one cell.

    A single bird routinely carries several parasite lineages at once, and a
    per-individual supplement reports that as a list in one cell. Harl et al
    2026's Table S1 has a "MalAvi lineages only" column reading
    ``lBUTBUT03, lBUBT3, hBUBT1`` -- three real associations for that bird, which
    the extractor previously kept as the single nonexistent lineage name
    ``BUTBUT03,LBUBT3,HBUBT1``. That is a false positive *and* three missed true
    positives, which is why co-infection cells cost twice.

    The split is deliberately strict: it fires only when the cell separates into
    **two or more parts that all look like lineage names**. One part that does not
    (a note, a morphospecies, ``MIX``, a bare number) means the cell is not a
    plain list and is left exactly as it was, for the curator to read. So this can
    only ever turn a cell that is unusable as a lineage name into the names it is
    actually made of.

    Returns a single-item list for an ordinary cell, so callers can treat every
    cell the same way.
    """
    if not value:
        return []
    text = str(value).strip()
    if not text:
        return []

    parts = [part.strip() for part in _COINFECTION_SPLIT.split(text)]
    parts = [part for part in parts if part]
    if len(parts) < 2 or not all(_LINEAGE_TOKEN.match(part) for part in parts):
        return [text]
    # Returned as printed. De-duplication happens on the *normalized* names at
    # emission, because ``lBUBT3`` and ``BUBT3`` are the same lineage printed two
    # ways and only normalize_lineage knows that.
    return parts


# A standalone integer inside a header, which in a wide layout is the slot index:
# "Lineage.1.Name" -> 1. Bounded by non-alphanumerics so that "cytb478" or
# "acc2no" are not mistaken for slots.
_SLOT_NUMBER = re.compile(r"(?<![a-z0-9])(\d{1,2})(?![a-z0-9])")


def _slot_of(cleaned_header: str) -> int:
    """Slot index of a header cell, or 0 when it is not part of a slot group."""
    match = _SLOT_NUMBER.search(cleaned_header)
    return int(match.group(1)) if match else 0


def _field_candidates(header: List[str]) -> Dict[int, Tuple[str, int, int]]:
    """Best MalAvi field for each column: ``column -> (field, end, synonym_len)``.

    **Rule 1 -- one column, several candidate fields: the rightmost match wins.**
    This exists because of the wide-format supplement in Fecchio et al 2023b,
    whose columns run ``Lineage.1.Genus | Lineage.1.Name | Lineage.1.Accession#``
    and where first-match-wins put the *parasite genus* ("PA", "HA", "LE") into
    ``lineage_name`` on all 1070 emitted records. These compound headers read
    ``Group.Index.Attribute``, so the attribute is the last token: "lineage 1
    genus" is a genus column, "lineage 1 accession" an accession column, and only
    "lineage 1 name" is the lineage itself. Scoring by where the matched synonym
    *ends* picks the attribute every time.
    """
    candidates: Dict[int, Tuple[str, int, int]] = {}
    for index, cell in enumerate(_clean(cell) for cell in header):
        if not cell:
            continue
        best: Optional[Tuple[str, int, int]] = None
        for field_name, synonyms in _COLUMN_SYNONYMS.items():
            for synonym in synonyms:
                position = cell.rfind(synonym)
                if position < 0:
                    continue
                end = position + len(synonym)
                # Rightmost match wins; a longer synonym breaks a tie.
                if field_name == "host_species" and _HOST_EXCLUSION.search(cell):
                    continue
                if best is None or (end, len(synonym)) > (best[1], best[2]):
                    best = (field_name, end, len(synonym))
        if best is not None:
            candidates[index] = best
    return candidates


def _rank_by_field(candidates: Dict[int, Tuple[str, int, int]],
                   columns: Optional[Set[int]] = None) -> Dict[str, List[int]]:
    """Invert ``_field_candidates``: ``field -> [column, ...]``, best first.

    A ranked list rather than one column, so the caller can fall through to the
    next candidate when the best one's values do not fit the field (see
    ``_column_values_fit``).

    **Rule 2 -- one field, several candidate columns: the most specific, then the
    leftmost, wins.** A longer matched synonym means a more specific header
    ("host species" beats a bare "host"), and among equals the first column is
    taken.

    ``columns`` restricts the ranking to one slot group, which is what stops
    ``Lineage.1.Name`` from being the only lineage column ever considered in a
    wide layout.
    """
    ranked: Dict[str, List[int]] = {}
    for field_name in _COLUMN_SYNONYMS:
        claiming = [(index, spec) for index, spec in candidates.items()
                    if spec[0] == field_name and (columns is None or index in columns)]
        if not claiming:
            continue
        claiming.sort(key=lambda item: (-item[1][2], item[0]))
        ranked[field_name] = [index for index, _ in claiming]
    return ranked


def _rank_columns(header: List[str]) -> Dict[str, List[int]]:
    """Ranked columns per field for a header row, ignoring slot structure."""
    return _rank_by_field(_field_candidates(header))


def _match_columns(header: List[str]) -> Dict[str, int]:
    """Top-ranked column per field, ignoring the values. Header-detection view."""
    return {field_name: columns[0]
            for field_name, columns in _rank_columns(header).items()}


# What a plausible value looks like, for the fields where a header alone has
# proved insufficient. A column is only accepted for one of these fields if
# enough of its actual values fit the shape.
#
# This exists because "Sequence number per host" (McNew et al 2021) matches the
# host synonym "host" perfectly well while holding the values 1, 2, 3 -- and the
# real host column in that file is called "Taxon", which matches nothing.
# Checking values rather than headers alone resolves both halves at once.
#
# The scalar name fields are anchored at BOTH ends. Without the closing anchor a
# cell only had to *begin* with a binomial, so "Turdus merula and other birds",
# "Turdus merula, lineage X" and "Turdus merula / Culex pipiens" were all accepted
# as a single host, and the vocabulary check -- which reads only the first token --
# could not catch any of them. A composite cell is not a scalar value; it is left
# for the curator instead of being silently truncated to its first name.
#
# What the anchored form still admits, deliberately: a trinomial (subspecies), a
# hyphenated epithet, and a trailing "sp."/"cf." qualifier, all of which are real
# ways a host is written in a table. host_names.canonical_host normalizes those.
_SCALAR_BINOMIAL = re.compile(
    r"^[A-Z][a-z]+[ _][a-z][a-z-]{1,}"           # Genus epithet
    r"(?:[ _](?:sp\.?|spp\.?|cf\.?|aff\.?|[a-z][a-z-]+))?$")   # optional 3rd token

_VALUE_SHAPE: Dict[str, "re.Pattern[str]"] = {
    # "Columbina talpacoti", "Cyanocompsa cyanoides" -- genus then epithet.
    "host_species": _SCALAR_BINOMIAL,
    # "CYACYA05", "AFR120", "T009", "G-016" -- a compact token, never a sentence.
    "lineage_name": re.compile(r"^[A-Za-z][A-Za-z0-9._-]{1,19}$"),
    # Vector binomials look exactly like host ones; the genus vocabulary is what
    # tells them apart (see _VALUE_VOCABULARY).
    "vector_species": _SCALAR_BINOMIAL,
    # "KM056426", "PX924994" -- an INSDC accession. Without this, a column of
    # running prose in a carved PDF block ("previously published sequences in
    # GenBank and thus", Schmid et al 2017a) is accepted as an accession column
    # because its header happens to say GenBank.
    "accession": re.compile(r"^[A-Z]{1,3}[_-]?\d{5,8}(\.\d+)?$"),
    # A run of nucleotides long enough to be a barcode fragment rather than a
    # word. The MalAvi frame is 479 bp and deposits are usually 470-479, but a
    # partial sequence is still worth resolving, so the floor is low; what it
    # excludes is everything that is not DNA.
    "sequence": re.compile(r"^[ACGTUNRYSWKMBDHV.-]{80,}$", re.IGNORECASE),
    # Counts are whole numbers and nothing else. This is what keeps a "positive"
    # synonym off a "% positive" or "No. positive pools (%)" column, whose values
    # read "11 (13)" -- and it is the guard that makes those two synonyms safe
    # enough to have at all, since inventing prevalence is a precision failure
    # the benchmark checks for explicitly.
    "number_tested": re.compile(r"^\d{1,7}$"),
    "number_found": re.compile(r"^\d{1,7}$"),
}

# Fraction of a column's sampled values that must fit the shape. The default is
# deliberately low: real tables carry blanks, "sp.", and the occasional malformed
# cell, and this check exists to reject columns that are wrong in kind, not
# imperfect. The counts are held to a much higher bar because their shape is so
# permissive -- a bare integer occurs in dozens of columns that are not counts
# (year, altitude, sample number), and a wrong count is a fabricated number
# rather than merely a name a curator would notice.
_VALUE_SHAPE_THRESHOLD = 0.30
_STRICT_SHAPE_THRESHOLD = 0.90
_STRICT_SHAPE_FIELDS = ("number_tested", "number_found")
_VALUE_SAMPLE_SIZE = 60


def _is_avian_binomial(value: str) -> bool:
    """Is this value a binomial whose genus is a known avian host genus?

    The vocabulary is the same MalAvi-derived gazetteer that ``hosts_geography``
    uses on prose (962 genera as of the 2026-03-23 release). Checking the genus
    rather than the full binomial keeps a *novel species in a known genus* --
    the usual shape of a new host record -- while rejecting a column that holds
    the wrong kind of organism entirely.

    That last case is not hypothetical. Kim & Tsuda 2012's Table 1 heads its
    mosquito column "Species", which matches the deliberately weak bare-"Species"
    host synonym, and its values (*Culex pipiens pallens*, *Aedes albopictus*)
    are perfectly well-formed binomials -- so neither the header guard nor the
    shape check can tell them from birds. Only the vocabulary can.
    """
    return _genus_in(value, _avian_host_genera())


def _is_vector_binomial(value: str) -> bool:
    """Is this value a binomial in a known arthropod-vector genus?

    Same idea as ``_is_avian_binomial``, against the gazetteer's vector genera
    (*Culex*, *Aedes*, *Culicoides*, …). Applied symmetrically so that a bird
    column can never be mistaken for a vector column either.
    """
    return _genus_in(value, _vector_genera())


def _genus_in(value: str, vocabulary: frozenset) -> bool:
    """Is the value's first token a genus in ``vocabulary``?

    An empty vocabulary means the gazetteer is unavailable, and the check goes
    vacuous rather than failing closed: without this, a missing data file would
    silently stop the extractor from finding any record at all.
    """
    if not vocabulary:
        return True
    genus = re.split(r"[ _]", value.strip(), maxsplit=1)[0]
    return genus.capitalize() in vocabulary


@lru_cache(maxsize=1)
def _avian_host_genera() -> frozenset:
    """Avian host genera from the bundled gazetteer; empty if it is missing."""
    return _gazetteer_set("genera")


@lru_cache(maxsize=1)
def _vector_genera() -> frozenset:
    """Arthropod-vector genera from the bundled gazetteer."""
    return _gazetteer_set("vector_genera")


def _gazetteer_set(key: str) -> frozenset:
    """One vocabulary from the packaged gazetteer, or empty if unavailable.

    Imported lazily so this module keeps working -- just less strictly -- for
    anyone using it without the packaged data file. An empty vocabulary makes
    the corresponding check vacuous rather than making it fail closed, because
    failing closed would silently stop extracting records altogether.
    """
    try:
        from .hosts_geography import load_gazetteer
        return frozenset(load_gazetteer().get(key, []))
    except Exception:
        return frozenset()


# Extra, vocabulary-based value tests layered on top of the shape patterns.
_VALUE_VOCABULARY = {"host_species": _is_avian_binomial,
                     "vector_species": _is_vector_binomial}


def _column_values_fit(matrix: List[List[Optional[str]]], start_row: int,
                       column: int, field_name: str) -> bool:
    """Does this column's data actually look like values of ``field_name``?

    Fields with no declared shape always pass -- the check is targeted, not a
    general-purpose filter.
    """
    pattern = _VALUE_SHAPE.get(field_name)
    if pattern is None:
        return True
    vocabulary = _VALUE_VOCABULARY.get(field_name)

    sampled: List[str] = []
    for row in matrix[start_row:]:
        if column >= len(row):
            continue
        value = _clean_value(row[column])
        if value:
            sampled.append(value)
        if len(sampled) >= _VALUE_SAMPLE_SIZE:
            break

    if not sampled:
        return False
    threshold = (_STRICT_SHAPE_THRESHOLD if field_name in _STRICT_SHAPE_FIELDS
                 else _VALUE_SHAPE_THRESHOLD)

    def fits(value: str) -> bool:
        if pattern.match(value) and (vocabulary is None or vocabulary(value)):
            return True
        # A lineage column may hold a co-infection -- several lineage names in
        # one cell. Such a cell is not shaped like a *single* lineage name, but
        # it is still a lineage column's value, and it is split into one row per
        # lineage on emission (see split_coinfection). Without this the shape
        # gate rejects the column outright and the whole table is lost.
        return field_name == "lineage_name" and len(split_coinfection(value)) > 1

    return sum(1 for value in sampled if fits(value)) / len(sampled) >= threshold


def _resolve_mapping(matrix: List[List[Optional[str]]], header_idx: int,
                     ranked: Dict[str, List[int]],
                     taken: Set[int]) -> Dict[str, int]:
    """Pick one column per field: best-ranked whose *values* fit the field.

    ``taken`` is shared across slot groups and mutated here, so no column is ever
    claimed by two fields. Falling through to the next candidate is what lets
    "Taxon" be found as the host column once "Sequence number per host" has been
    rejected on its values.
    """
    mapping: Dict[str, int] = {}
    for field_name, columns in ranked.items():
        for column in columns:
            if column in taken:
                continue
            if not _column_values_fit(matrix, header_idx + 1, column, field_name):
                continue
            mapping[field_name] = column
            taken.add(column)
            break
    return mapping


def _slot_groups(header: List[str],
                 candidates: Dict[int, Tuple[str, int, int]]) -> Dict[int, Set[int]]:
    """Partition the mapped columns into a base group (0) and any slot groups.

    A wide table repeats a group of columns once per co-infecting lineage:
    Fecchio et al 2023b runs ``Lineage.1.Genus | Lineage.1.Name |
    Lineage.1.Accession#`` through ``Lineage.5.*`` alongside single-valued
    columns for the host, the country and the date. Only slot 1 was ever read,
    which capped that paper's recall at 71% no matter how well the parser worked.

    A numbered header is only believed to be a slot when the layout is genuinely
    repeated: at least two indices whose **field signatures match**, meaning the
    two groups resolve the same set of MalAvi fields, one of which is the lineage.

    Requiring the signature to repeat -- not merely the number -- is what stops a
    number that means something else entirely from expanding a row into several
    records. ``PCR 1 lineage | PCR 2 lineage`` are technical replicates of one
    sample and ``Site 1 lineage | Site 2 lineage`` are two localities; read as
    co-infection slots, both turn one row into two records, and the replicate case
    emits the *same* lineage twice. A real slot layout looks like Fecchio et al
    2023b's, where each index carries a genus, a name and an accession, so the
    signature ``{parasite_genus, lineage_name, accession}`` repeats verbatim.
    """
    cleaned = [_clean(cell) for cell in header]
    by_slot: Dict[int, Set[int]] = {}
    for column in candidates:
        slot = _slot_of(cleaned[column]) if column < len(cleaned) else 0
        by_slot.setdefault(slot, set()).add(column)

    numbered = {slot: columns for slot, columns in by_slot.items() if slot > 0}
    if len(numbered) < 2:
        return {0: set(candidates)}

    # The signature of a group is the set of fields its columns resolve.
    signatures: Dict[int, frozenset] = {
        slot: frozenset(candidates[column][0] for column in columns)
        for slot, columns in numbered.items()
    }
    # A slot layout needs one signature, containing the lineage, shared by two or
    # more indices. A signature of *only* the lineage is not enough on its own:
    # that is exactly the shape "PCR 1 lineage | PCR 2 lineage" has, and it says
    # nothing about whether the repetition means co-infection.
    repeated = {
        signature for signature in signatures.values()
        if "lineage_name" in signature and len(signature) >= 2
        and sum(1 for other in signatures.values() if other == signature) >= 2
    }
    if not repeated:
        return {0: set(candidates)}

    # Only the groups that actually match the repeated signature are slots;
    # anything else numbered incidentally stays in the base group.
    slots: Dict[int, Set[int]] = {0: set(by_slot.get(0, set()))}
    for slot, columns in numbered.items():
        if signatures[slot] in repeated:
            slots[slot] = columns
        else:
            slots[0] |= columns
    return slots


@dataclass
class BlockRows:
    """What one table block yielded, split by MalAvi table.

    Kept apart because they are governed by different rules: an association
    (record or vector) must be complete to exist at all, whereas an accession is
    a fact the paper states on its own.
    """

    records: List[Dict[str, Optional[str]]] = field(default_factory=list)
    vectors: List[Dict[str, Optional[str]]] = field(default_factory=list)
    accessions: List[str] = field(default_factory=list)


# How far down a block to look for the header row.
_HEADER_SEARCH_ROWS = 15


def _best_header_row(matrix: List[List[Optional[str]]]) -> Optional[int]:
    """Which row near the top of a block is its header?

    Every row in the search window is scored and the best one wins. Taking the
    *first* row that mapped two fields -- the previous rule -- locks onto whatever
    matched earliest, and a caption, a title, a units row or the upper tier of a
    two-tier header all match readily: "Table 2. Lineages and hosts of..." resolves
    both required fields without being a header at all. The real header below it
    was then never considered, because the search had already stopped.

    The score is deterministic and made only of things that are checkable:

      * how many distinct MalAvi fields the row resolves *and* whose columns hold
        values of the right shape below it (the only evidence that matters), then
      * how many fields it resolves at all, then
      * how many of its cells are short enough to be headers rather than prose.

    Earlier rows win ties, so a genuine header is not passed over for a data row
    that happens to score the same.
    """
    best: Optional[Tuple[Tuple[int, int, int], int]] = None
    for index, row in enumerate(matrix[:_HEADER_SEARCH_ROWS]):
        header = [str(cell) if cell is not None else "" for cell in row]
        ranked = _rank_columns(header)
        if len(ranked) < 2:
            continue

        # Fields whose best column actually holds values of the right kind. This
        # is what separates a header from a sentence that happens to contain the
        # words "host" and "lineage": under a real header the column below it holds
        # binomials or lineage codes, and under a caption it holds nothing of the
        # sort.
        supported = sum(
            1 for field_name, columns in ranked.items()
            if any(_column_values_fit(matrix, index + 1, column, field_name)
                   for column in columns))
        # Header cells are labels. A row of prose is not a header even if it names
        # the right fields.
        label_like = sum(1 for cell in header if cell and len(cell.strip()) <= 40)
        score = (supported, len(ranked), label_like)
        if best is None or score > best[0]:
            best = (score, index)

    return None if best is None else best[1]


def _rows_from_matrix(
    matrix: List[List[Optional[str]]],
) -> Tuple[BlockRows, Dict[str, str]]:
    """Find the header row, map columns, and emit structured rows.

    Returns (block_rows, columns_detected) where columns_detected maps each
    MalAvi field to the actual header text it was matched from.
    """
    header_idx = _best_header_row(matrix)
    if header_idx is None:
        return BlockRows(), {}

    header = [str(cell) if cell is not None else "" for cell in matrix[header_idx]]
    candidates = _field_candidates(header)
    groups = _slot_groups(header, candidates)

    # The base group's columns apply to every record from a row (the host, the
    # locality); each slot group contributes one record's worth on top.
    taken: Set[int] = set()
    base_mapping = _resolve_mapping(
        matrix, header_idx, _rank_by_field(candidates, groups.get(0, set())), taken)
    slot_mappings = {
        slot: _resolve_mapping(
            matrix, header_idx, _rank_by_field(candidates, columns), taken)
        for slot, columns in sorted(groups.items()) if slot > 0
    }

    # Every field this block resolved, across the base group and all slots. A
    # block that can supply neither a host association nor a vector association
    # nor a column of accessions holds nothing MalAvi wants.
    resolved_fields = set(base_mapping)
    for mapping in slot_mappings.values():
        resolved_fields.update(mapping)
    has_records = all(f in resolved_fields for f in _REQUIRED_RECORD_FIELDS)
    has_vectors = all(f in resolved_fields for f in _REQUIRED_VECTOR_FIELDS)
    has_accessions = "accession" in resolved_fields
    if not (has_records or has_vectors or has_accessions):
        return BlockRows(), {}

    columns_detected = {
        field_name: (header[column] if column < len(header) else "")
        for mapping in [base_mapping, *slot_mappings.values()]
        for field_name, column in mapping.items()
    }

    def read(row: List[Optional[str]], mapping: Dict[str, int]) -> Dict[str, Optional[str]]:
        """Pull one mapping's fields out of a data row."""
        record: Dict[str, Optional[str]] = {}
        for field_name, column in mapping.items():
            value = _clean_value(row[column]) if column < len(row) else None
            # A category word is not a name. Dropping it here (rather than at the
            # column level) is what a per-row check buys: the column is a real
            # lineage column, and only this one cell says "MIX".
            if _is_reserved(value) and field_name in _NAMED_FIELDS:
                value = None
            if field_name == "lineage_name":
                value = normalize_lineage(value)
            record[field_name] = value
        return record

    def emit(out: BlockRows, candidate: Dict[str, Optional[str]]) -> None:
        """File a parsed row, expanding a co-infection cell into one row each.

        A cell naming several lineages is several associations sharing one host,
        locality and sample -- so every other field is copied unchanged onto each
        row. Cells that are not a plain list of lineage names pass through as a
        single row exactly as before (see ``split_coinfection``).
        """
        lineages = split_coinfection(candidate.get("lineage_name"))
        if len(lineages) < 2:
            sort_into(out, candidate)
            return
        emitted: Set[str] = set()
        for lineage in lineages:
            name = normalize_lineage(lineage)
            # One cell can print the same lineage twice (``lBUBT3, hBUBT1,
            # BUBT3``); that is one association, not two.
            if not name or name in emitted:
                continue
            emitted.add(name)
            expanded = dict(candidate)
            expanded["lineage_name"] = name
            # Keep what the cell said, so a curator can see the co-infection this
            # row was split out of rather than a name with no context.
            expanded["lineage_name_source"] = candidate.get("lineage_name")
            sort_into(out, expanded)

    def sort_into(out: BlockRows, candidate: Dict[str, Optional[str]]) -> None:
        """File one parsed row under records, vectors and/or accessions.

        A row can be more than one of these at once: Perrin et al 2026's Table S4
        row is both a vector association and an accession. Accessions are
        collected whether or not the row is an association, because an accession
        is a *set* the paper reports, not a claim about who hosted what.

        The host and vector tests are independent (two ``if``s, not ``if``/
        ``elif``). A table can legitimately report a lineage, the bird it came
        from and the mosquito it was also found in on one row; under ``elif`` the
        vector association was silently dropped whenever a host was present,
        which contradicted this docstring.
        """
        if all(candidate.get(f) for f in _REQUIRED_RECORD_FIELDS):
            out.records.append(candidate)
        if all(candidate.get(f) for f in _REQUIRED_VECTOR_FIELDS):
            out.vectors.append(candidate)
        if candidate.get("accession"):
            out.accessions.append(candidate["accession"])

    out = BlockRows()
    for row in matrix[header_idx + 1:]:
        base = read(row, base_mapping)
        if not slot_mappings:
            emit(out, base)
            continue
        # Wide layout: one record per occupied slot, each carrying the base
        # columns. A slot with no lineage is an unused co-infection column.
        for mapping in slot_mappings.values():
            slot_values = read(row, mapping)
            if not slot_values.get("lineage_name"):
                continue
            record = dict(base)
            record.update({k: v for k, v in slot_values.items() if v is not None})
            emit(out, record)
    return out, columns_detected


# A block is one named table: (name, matrix). The name is the sheet name, the
# Word table's caption, or the file name for a CSV -- whatever a human would call
# it, which is also what the metadata-block check reads.
Block = Tuple[str, List[List[Optional[str]]]]


def _read_csv(path: Path, delimiter: str) -> List[Block]:
    """A delimited file is a single block, named for the file."""
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        matrix: List[List[Optional[str]]] = [row for row in csv.reader(fh, delimiter=delimiter)]
    return [(path.stem, matrix)]


def _read_xlsx(path: Path) -> List[Block]:
    """One block per worksheet, named by the sheet name."""
    import openpyxl  # optional 'tables' extra

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks: List[Block] = []
    for worksheet in workbook.worksheets:
        matrix = [[None if cell is None else str(cell) for cell in row]
                  for row in worksheet.iter_rows(values_only=True)]
        blocks.append((worksheet.title, matrix))
    workbook.close()
    return blocks


def _read_docx(path: Path) -> List[Block]:
    """One block per Word table, named by the caption paragraph above it.

    Word documents carry no table names, so the body is walked in document order
    and the most recent non-empty paragraph is used as the caption. That is where
    "Table S2. Vertebrate host-parasite interactions detected in the study" lives,
    which both names the block and lets the metadata check see it.
    """
    import docx  # python-docx, optional 'tables' extra
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    blocks: List[Block] = []
    caption = ""
    table_number = 0

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text.strip()
            if text:
                caption = text
        elif child.tag.endswith("}tbl"):
            table_number += 1
            table = Table(child, document)
            matrix: List[List[Optional[str]]] = [
                [cell.text for cell in row.cells] for row in table.rows]
            # Keep the caption short: it is a label, not the table's abstract.
            name = caption[:120] if caption else f"table {table_number}"
            blocks.append((name, matrix))
            caption = ""  # a caption belongs to one table only

    return blocks


def extract_blocks(blocks: List[Block], source: Path) -> ExtractedRows:
    """Parse a list of named tables into structured rows + flattened text.

    Every block is parsed independently and their rows unioned, so a document's
    second and later tables contribute too. Blocks whose name marks them as
    documentation are skipped for *records* but still contribute their text,
    because accessions are frequently listed in a legend.

    This is the single path every source goes down -- worksheet, Word table, CSV
    or a table carved out of a PDF -- so a fix to column matching benefits all of
    them at once.
    """
    combined = BlockRows()
    merged_columns: Dict[str, str] = {}
    block_reports: List[Dict[str, object]] = []
    text_parts: List[str] = []

    for name, matrix in blocks:
        # The flattened text of every block, including skipped ones, still feeds
        # the accession miner.
        text_parts.append(
            "\n".join(" ".join("" if c is None else str(c) for c in row)
                      for row in matrix))

        if _METADATA_BLOCK.search(name):
            block_reports.append({"block": name, "rows": 0,
                                  "skipped": "documentation block, not data"})
            continue

        parsed, columns = _rows_from_matrix(matrix)
        # Which block each row came from, so a deduplicated row can still say
        # where it was seen.
        for row in parsed.records + parsed.vectors:
            row.setdefault("source_block", name)
        combined.records.extend(parsed.records)
        combined.vectors.extend(parsed.vectors)
        combined.accessions.extend(parsed.accessions)
        block_reports.append({"block": name, "rows": len(parsed.records),
                              "vector_rows": len(parsed.vectors),
                              "columns_detected": columns})
        # First productive block wins a given field name in the merged summary;
        # the per-block detail carries the rest. This summary is a convenience
        # only, and it is genuinely lossy -- a wide slot layout resolves the same
        # field in several groups, and Fecchio et al 2023b's summary therefore
        # reads "Lineage.5.Name". Read `blocks` for what each block actually did.
        for field_name, header in columns.items():
            merged_columns.setdefault(field_name, header)

    records, n_duplicate_records = _deduplicate(combined.records)
    vectors, n_duplicate_vectors = _deduplicate(combined.vectors)

    return ExtractedRows(source=source, rows=records,
                         vectors=vectors,
                         accessions=sorted(set(combined.accessions)),
                         text="\n".join(text_parts),
                         columns_detected=merged_columns, blocks=block_reports,
                         n_duplicate_rows=n_duplicate_records + n_duplicate_vectors)


# Fields that identify a record. Two rows agreeing on all of these are the same
# claim, however many surfaces it was read through.
_IDENTITY_FIELDS = ("lineage_name", "host_species", "vector_species", "country",
                    "site", "number_tested", "number_found")


def _deduplicate(rows: List[Dict[str, Optional[str]]],
                 ) -> Tuple[List[Dict[str, Optional[str]]], int]:
    """Collapse identical rows, keeping every place each was seen.

    The same printed table reaches this module through more than one door -- a
    supplement can ship as both a PDF and a spreadsheet, a multi-page table
    repeats its header on each page, and a paper's own table is carved out of the
    layout text as well as read from the file. Harl et al 2026's Table 1 emits
    ``CIAE11 x Circus aeruginosus`` twice for that reason.

    Without this, one association looks better supported merely because it was
    extracted through several surfaces, and the curator reviews the same claim
    repeatedly. Collapsing is safe *because* the identity includes the counts: two
    rows that agree on lineage, host, place and prevalence are not two
    observations, they are one observation read twice.

    Returns (unique rows, number collapsed). The surviving row gains
    ``source_blocks``: every block the claim was seen in.
    """
    unique: Dict[tuple, Dict[str, Optional[str]]] = {}
    duplicates = 0
    for row in rows:
        key = tuple((row.get(f) or "").strip().lower() if isinstance(row.get(f), str)
                    else row.get(f) for f in _IDENTITY_FIELDS)
        existing = unique.get(key)
        if existing is None:
            row["source_blocks"] = [row["source_block"]] if row.get("source_block") else []
            unique[key] = row
            continue
        duplicates += 1
        block = row.get("source_block")
        blocks_seen = existing.setdefault("source_blocks", [])
        if block and block not in blocks_seen:
            blocks_seen.append(block)
    return list(unique.values()), duplicates


def extract_table_file(path: str | Path) -> ExtractedRows:
    """Extract structured rows + text from a supplementary table file."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()
    if suffix in (".csv",):
        blocks = _read_csv(path, ",")
    elif suffix in (".tsv", ".tab"):
        blocks = _read_csv(path, "\t")
    elif suffix in (".xlsx", ".xlsm"):
        blocks = _read_xlsx(path)
    elif suffix == ".xls":
        # openpyxl cannot read the legacy binary .xls format at all, so routing
        # these here only produced a confusing "corrupt file" error. Saying so
        # plainly is more useful: a curator can re-save the file as .xlsx in
        # seconds, whereas a silent failure loses the whole supplement.
        raise ValueError(
            f"Legacy binary .xls is not supported ({path.name}); re-save it as "
            ".xlsx (or export to .csv) and run again")
    elif suffix in (".docx",):
        blocks = _read_docx(path)
    else:
        raise ValueError(f"Unsupported supplementary file type: {suffix}")

    return extract_blocks(blocks, source=path)


def extract_pdf_tables(document) -> ExtractedRows:
    """Structured rows from the tables printed inside a PDF.

    Takes an already-parsed ``pdf_extract.ExtractedDocument`` (so the PDF is read
    once, not once per consumer) and runs its carved layout tables through the
    same block parser as a spreadsheet.

    Six of the ten papers in the ground-truth corpus ship no supplementary
    spreadsheet at all, and pdfplumber finds no ruled tables in four of them
    because their tables have no ruling lines. For those papers this is the only
    route by which a lineage x host association can be read at all.

    Only the ``text`` of blocks is *not* returned here: the PDF's full text
    surface already reaches the accession miner via ``mining_text()``, so
    returning it again would double-count.
    """
    rows = extract_blocks(document.layout_tables(), source=document.path)
    rows.text = ""
    return rows
